"""Referat matnidan tayyor .docx fayl yaratish."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

OUTPUT_DIR = "generated_files"
os.makedirs(OUTPUT_DIR, exist_ok=True)


def create_referat_docx(mavzu: str, sections: dict, user_id: int) -> str:
    doc = Document()

    # Umumiy shrift sozlamalari
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)

    section = doc.sections[0]
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    # --- Titul sahifa ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("REFERAT")
    run.bold = True
    run.font.size = Pt(20)

    for _ in range(6):
        doc.add_paragraph()

    mavzu_p = doc.add_paragraph()
    mavzu_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mavzu_run = mavzu_p.add_run(f'Mavzu: "{mavzu}"')
    mavzu_run.bold = True
    mavzu_run.font.size = Pt(16)

    doc.add_page_break()

    # --- Mundarija (oddiy) ---
    toc_title = doc.add_heading("MUNDARIJA", level=1)
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for item in ["Kirish", "Asosiy qism", "Xulosa", "Foydalanilgan adabiyotlar"]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_page_break()

    # --- Asosiy bo'limlar ---
    def add_section(heading_text, body_text):
        h = doc.add_heading(heading_text, level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for paragraph in body_text.split("\n"):
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            p = doc.add_paragraph(paragraph)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.25)
            p.paragraph_format.line_spacing = 1.5

    add_section("KIRISH", sections.get("kirish", ""))
    add_section("ASOSIY QISM", sections.get("asosiy_qism", ""))
    add_section("XULOSA", sections.get("xulosa", ""))

    # Adabiyotlar ro'yxati alohida (bullet emas, oddiy qatorlar)
    doc.add_heading("FOYDALANILGAN ADABIYOTLAR", level=1)
    for line in sections.get("adabiyotlar", "").split("\n"):
        line = line.strip()
        if line:
            doc.add_paragraph(line)

    filename = f"referat_{user_id}.docx"
    filepath = os.path.join(OUTPUT_DIR, filename)
    doc.save(filepath)
    return filepath
