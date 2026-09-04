"""Taqdimot tuzilmasidan tayyor .pptx fayl yaratish."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUTPUT_DIR = "generated_files"
os.makedirs(OUTPUT_DIR, exist_ok=True)

PRIMARY_COLOR = RGBColor(0x1F, 0x4E, 0x79)
ACCENT_COLOR = RGBColor(0x2E, 0x86, 0xC1)


def create_presentation_pptx(outline: dict, user_id: int) -> str:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --- Sarlavha slaydi ---
    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = outline.get("title", "Taqdimot")
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(40)
    slide.shapes.title.text_frame.paragraphs[0].font.bold = True
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR

    if len(slide.placeholders) > 1:
        subtitle = outline.get("subtitle", "")
        slide.placeholders[1].text = subtitle

    # --- Kontent slaydlari ---
    bullet_layout = prs.slide_layouts[1]
    for slide_data in outline.get("slides", []):
        s = prs.slides.add_slide(bullet_layout)
        s.shapes.title.text = slide_data.get("title", "")
        s.shapes.title.text_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR
        s.shapes.title.text_frame.paragraphs[0].font.bold = True

        body = s.placeholders[1].text_frame
        body.clear()
        bullets = slide_data.get("bullets", [])
        for i, bullet in enumerate(bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = str(bullet)
            p.level = 0
            p.font.size = Pt(22)

    filename = f"prezentatsiya_{user_id}.pptx"
    filepath = os.path.join(OUTPUT_DIR, filename)
    prs.save(filepath)
    return filepath
